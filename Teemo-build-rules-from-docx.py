"""Convert Teemo's DOCX reply rules into the Markdown consumed by the app."""

from pathlib import Path
import argparse

from docx import Document


def front_matter(document):
    metadata = {
        "name": "Teemo私人助理回复规则",
        "description": "Teemo个人AI助理的默认回复、任务判断、工具调用与上下文管理规则",
        "icon": "🤖",
    }
    if document.tables:
        for row in document.tables[0].rows:
            text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if ":" not in text:
                continue
            key, value = text.split(":", 1)
            if key.strip() in metadata and value.strip():
                metadata[key.strip()] = value.strip()
    return [
        "---",
        f"name: {metadata['name']}",
        f"description: {metadata['description']}",
        f"icon: {metadata['icon']}",
        "---",
        "",
    ]


def paragraph_to_markdown(paragraph):
    text = paragraph.text.strip()
    if not text:
        return []

    style = paragraph.style.name if paragraph.style else ""
    if style == "Heading 1":
        return [f"# {text}", ""]
    if style == "Heading 2":
        return [f"## {text}", ""]
    if style == "Heading 3":
        return [f"### {text}", ""]

    properties = paragraph._p.pPr
    numbered = properties is not None and properties.numPr is not None
    if numbered:
        return [f"- {text.replace(chr(10), '  ' + chr(10))}"]

    # Preserve intentional line breaks such as the DesignHub call chain.
    return [text.replace("\n", "  \n"), ""]


def convert(input_path, output_path):
    document = Document(input_path)
    lines = front_matter(document)
    previous_was_list = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        properties = paragraph._p.pPr
        is_list = bool(text and properties is not None and properties.numPr is not None)
        if previous_was_list and not is_list and lines and lines[-1].strip():
            lines.append("")
        lines.extend(paragraph_to_markdown(paragraph))
        previous_was_list = is_list

    while lines and not lines[-1].strip():
        lines.pop()
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    convert(args.input, args.output)


if __name__ == "__main__":
    main()
